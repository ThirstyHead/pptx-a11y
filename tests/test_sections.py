"""Tests for section naming rules (WCAG 2.4.2 / 1.3.1)."""
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.oxml import parse_xml
from pptx_a11y.rules import DefaultSectionNameRule, DuplicateSectionNameRule, AuditContext


def _add_sections_to_prs(prs: PresentationType, names: list[str]):
    # Add section list to prs._element
    section_xml_items = []
    for idx, name in enumerate(names, start=1):
        section_xml_items.append(f'<p:section xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" name="{name}" id="{{11111111-2222-3333-4444-55555555555{idx}}}"><p:sldIdLst/></p:section>')
    section_lst_xml = f'<p:sectionLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">{"".join(section_xml_items)}</p:sectionLst>'
    prs._element.append(parse_xml(section_lst_xml))


def test_default_section_name_detected():
    prs = Presentation()
    _add_sections_to_prs(prs, ["Default Section", "Financials"])

    rule = DefaultSectionNameRule()
    findings = rule.check(prs, AuditContext("sections.pptx"))
    assert len(findings) == 1
    assert findings[0].rule_id == "section-name-default"
    assert findings[0].sc == "2.4.2"
    assert findings[0].fixable is True


def test_duplicate_section_name_detected():
    prs = Presentation()
    _add_sections_to_prs(prs, ["Financials", "Financials"])

    rule = DuplicateSectionNameRule()
    findings = rule.check(prs, AuditContext("sections.pptx"))
    assert len(findings) == 1
    assert findings[0].rule_id == "section-name-duplicate"
    assert findings[0].sc == "2.4.2"


def test_section_remediation(tmp_path):
    prs = Presentation()
    _add_sections_to_prs(prs, ["Default Section", "Default Section"])
    in_path = tmp_path / "in_sec.pptx"
    out_path = tmp_path / "out_sec.pptx"
    prs.save(str(in_path))

    from pptx_a11y.remediate import remediate_presentation
    fixes = remediate_presentation(in_path, out_path)
    assert fixes["sections_renamed"] >= 2

    # Audit remediated file
    from pptx_a11y.audit import audit_file
    res = audit_file(out_path)
    rule_ids = {f["rule_id"] for f in res["findings"]}
    assert "section-name-default" not in rule_ids
    assert "section-name-duplicate" not in rule_ids
