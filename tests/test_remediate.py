"""Tests for deterministic presentation remediation."""
from pathlib import Path
from pptx_a11y.audit import audit_file
from pptx_a11y.remediate import remediate_presentation

FIXTURES = Path(__file__).parent / "fixtures"


def test_remediation_reduces_blocking_violations(tmp_path: Path):
    in_file = FIXTURES / "violations.pptx"
    out_file = tmp_path / "violations_fixed.pptx"

    fixes = remediate_presentation(in_file, out_file)
    assert fixes["title_added"] > 0
    assert fixes["table_headers_set"] > 0

    # Audit the remediated deck
    res_before = audit_file(in_file)
    res_after = audit_file(out_file)

    assert res_after["summary"]["blocking"] < res_before["summary"]["blocking"]
    rule_ids_after = {f["rule_id"] for f in res_after["findings"]}
    assert "title-missing" not in rule_ids_after
    assert "table-header-missing" not in rule_ids_after


def test_original_file_immutability_and_sha256(tmp_path: Path):
    import hashlib
    import pytest

    in_file = FIXTURES / "violations.pptx"
    out_file = tmp_path / "violations_fixed.pptx"

    sha_before = hashlib.sha256(in_file.read_bytes()).hexdigest()
    fixes = remediate_presentation(in_file, out_file)
    sha_after = hashlib.sha256(in_file.read_bytes()).hexdigest()

    assert sha_before == sha_after
    assert fixes["original_sha256"] == sha_before
    assert fixes["original_file_immutable"] is True

    # In-place remediation must be rejected to prevent corrupting or altering the original file
    with pytest.raises(ValueError, match="Cannot remediate in place"):
        remediate_presentation(in_file, in_file)


def test_remediation_does_not_strip_passwords_or_security(tmp_path: Path):
    from pptx import Presentation
    from pptx.oxml import parse_xml

    # Create a dummy deck with modifyVerifier
    in_file = tmp_path / "protected.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    mv = parse_xml('<p:modifyVerifier xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" cryptProviderType="rsaAES" cryptAlgorithmClass="hash"/>')
    prs._element.append(mv)
    prs.save(str(in_file))

    out_file = tmp_path / "remediated.pptx"
    fixes = remediate_presentation(in_file, out_file)

    # Remediation engine must NOT strip passwords / security verifiers
    assert "passwords_stripped" not in fixes
    assert "restricted_access_removed" not in fixes

    prs_out = Presentation(str(out_file))
    verifiers = prs_out._element.xpath(".//p:modifyVerifier | .//*[local-name()='modifyVerifier']")
    assert len(verifiers) == 1
