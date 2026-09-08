"""Tests for interactive CLI triage mode."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx_a11y.triage import run_interactive_triage
from pptx_a11y.audit import audit_file


def test_interactive_triage_sets_alt_text_and_titles(tmp_path: Path):
    in_deck = tmp_path / "triage_input.pptx"
    out_deck = tmp_path / "triage_output.pptx"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title:
        slide.shapes.title.text = "Financial Overview"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    shape.name = "Diagram 1"
    prs.save(str(in_deck))

    # Simulate user answering alt text prompt
    inputs = ["Detailed revenue growth diagram"]
    logs = []

    def mock_input(prompt):
        return inputs.pop(0)

    def mock_print(*args):
        logs.append(" ".join(str(a) for a in args))

    count = run_interactive_triage(
        in_path=in_deck,
        out_path=out_deck,
        input_func=mock_input,
        print_func=mock_print,
    )

    assert count == 1
    assert out_deck.exists()

    # Re-audit the triaged file: image-alt-missing should now be resolved
    res = audit_file(out_deck)
    rule_ids = {f["rule_id"] for f in res["findings"]}
    assert "image-alt-missing" not in rule_ids


def test_interactive_triage_marks_decorative(tmp_path: Path):
    in_deck = tmp_path / "triage_dec_input.pptx"
    out_deck = tmp_path / "triage_dec_output.pptx"

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title:
        slide.shapes.title.text = "Slide Title"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    shape.name = "Background Box"
    prs.save(str(in_deck))

    inputs = ["d"]

    count = run_interactive_triage(
        in_path=in_deck,
        out_path=out_deck,
        input_func=lambda p: inputs.pop(0),
        print_func=lambda *a: None,
    )

    assert count == 1
    res = audit_file(out_deck)
    rule_ids = {f["rule_id"] for f in res["findings"]}
    assert "image-alt-missing" not in rule_ids
