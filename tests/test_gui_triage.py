"""Tests for GUI interactive TriageDialog."""
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx_a11y.audit import audit_file
from pptx_a11y.gui.triage_dialog import TriageDialog


def test_triage_dialog_flow(qtbot, tmp_path: Path):
    deck = tmp_path / "test_triage.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    if slide.shapes.title:
        slide.shapes.title.text = "Intro"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    shape.name = "Flowchart 1"
    prs.save(str(deck))

    audit_res = audit_file(deck)
    findings = [f for f in audit_res["findings"] if f["rule_id"] in ("image-alt-missing", "slide-title-missing")]
    assert len(findings) >= 1

    dialog = TriageDialog(deck, findings)
    qtbot.addWidget(dialog)

    # Supply text and apply
    dialog.txt_input.setText("Company organizational flowchart")
    dialog.apply_current()

    # Re-audit: image-alt-missing should be resolved
    res2 = audit_file(deck)
    rule_ids = {f["rule_id"] for f in res2["findings"]}
    assert "image-alt-missing" not in rule_ids
