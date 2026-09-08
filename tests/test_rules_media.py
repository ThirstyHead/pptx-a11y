"""Tests for media subtitles rule (WCAG 1.2.2)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx_a11y.rules import MediaSubtitlesRule, AuditContext


def test_media_subtitles_rule_flags_unsubtitled_video(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add a media/video placeholder element in oxml
    # Simulate a video shape with <a:videoFile>
    tx = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
    # inject a:videoFile in oxml to simulate media shape
    from pptx.oxml import parse_xml
    video_elem = parse_xml('<a:videoFile xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:link="rId1"/>')
    tx._element.spPr.append(video_elem)

    rule = MediaSubtitlesRule()
    ctx = AuditContext(source_name="media_test.pptx")
    findings = rule.check(prs, ctx)

    assert len(findings) == 1
    f = findings[0]
    assert f.rule_id == "media-subtitles-missing"
    assert f.sc == "1.2.2"
    assert f.severity == "critical"
    assert f.why_unfixable is not None
    assert len(f.manual_steps) >= 3
