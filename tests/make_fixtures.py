"""Generate synthetic .pptx test files for unit and integration testing."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

FIXTURES_DIR = Path(__file__).parent / "fixtures"


def make_clean_deck(path: Path):
    prs = Presentation()
    prs.core_properties.title = "Accessible Quarterly Business Review"
    prs.core_properties.language = "en-US"

    # Slide 1: Title Slide (Semantic title layout)
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    if slide1.shapes.title:
        slide1.shapes.title.text = "Q4 Accessible Presentation"
    if len(slide1.placeholders) > 1:
        subtitle = slide1.placeholders[1]
        tf = getattr(subtitle, "text_frame", None)
        if tf:
            tf.text = "Executive Summary"

    # Slide 2: Title & Content with Table featuring explicit header
    bullet_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_layout)
    if slide2.shapes.title:
        slide2.shapes.title.text = "Financial Highlights"

    # Tag all text runs with language
    for slide in prs.slides:
        for shape in slide.shapes:
            tf = getattr(shape, "text_frame", None)
            if tf:
                for p in tf.paragraphs:
                    for r in p.runs:
                        rPr = r._r.get_or_add_rPr()
                        rPr.set("lang", "en-US")

    # Add a clean table
    rows, cols = 3, 2
    table_shape = slide2.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(6), Inches(2))
    table = table_shape.table
    # Set headers
    table.cell(0, 0).text = "Quarter"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "Q3"
    table.cell(1, 1).text = "$1.2M"
    table.cell(2, 0).text = "Q4"
    table.cell(2, 1).text = "$1.5M"

    # Ensure header row attribute is explicitly set on a:tblPr
    tblPr = table_shape._element.xpath(".//a:tblPr")
    if tblPr:
        tblPr[0].set("firstRow", "1")

    prs.save(str(path))


def make_violations_deck(path: Path):
    prs = Presentation()
    # Violation: Missing presentation title in core properties
    prs.core_properties.title = ""

    # Slide 1: Blank layout with floating text box pretending to be a title (Violates 1.3.1, 2.4.2)
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Fake Title Box (No Semantic Layout)"

    # Add unlabelled shape (Violates 1.1.1)
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(2), Inches(2))
    # No alt text set!

    # Slide 2: Table missing header row declaration
    slide2 = prs.slides.add_slide(blank_layout)
    table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(1.5))
    table = table_shape.table
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    # Ensure firstRow is false or omitted
    tblPr = table_shape._element.xpath(".//a:tblPr")
    if tblPr:
        tblPr[0].set("firstRow", "0")

    prs.save(str(path))


if __name__ == "__main__":
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    make_clean_deck(FIXTURES_DIR / "clean.pptx")
    make_violations_deck(FIXTURES_DIR / "violations.pptx")
    print(f"Generated fixtures in {FIXTURES_DIR}")
