"""Tests for reading order analyzer and spatial sorting (WCAG 1.3.2)."""
from pptx import Presentation
from pptx.util import Inches
from pptx_a11y.reading_order import check_slide_reading_order, reorder_slide_shapes
from pptx_a11y.rules import ReadingOrderRule, AuditContext


def test_reading_order_inversion_detected():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add shape 1 at bottom of slide
    s1 = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    s1.name = "BottomBox"

    # Add shape 2 at top of slide (placed later in DOM, but visually higher!)
    s2 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    s2.name = "TopBox"

    inverted = check_slide_reading_order(slide)
    assert inverted is True

    rule = ReadingOrderRule()
    findings = rule.check(prs, AuditContext("test_order.pptx"))
    assert len(findings) == 1
    assert findings[0].rule_id == "reading-order-inverted"
    assert findings[0].sc == "1.3.2"


def test_reading_order_reorder_shapes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    s1 = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    s1.name = "BottomBox"
    s2 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    s2.name = "TopBox"

    reorder_slide_shapes(slide)
    # After reordering, visual top shape should come first in DOM tree
    content_shapes = [s.name for s in slide.shapes]
    assert content_shapes[0] == "TopBox"
    assert content_shapes[1] == "BottomBox"
    assert check_slide_reading_order(slide) is False
