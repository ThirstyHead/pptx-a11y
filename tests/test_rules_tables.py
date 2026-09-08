"""Tests for table merged cells rule (WCAG 1.3.1)."""
from pptx import Presentation
from pptx.util import Inches
from pptx_a11y.rules import TableMergedCellsRule, AuditContext


def test_table_merged_cells_detected():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))

    # Simulate merged cells in oxml
    cell_elem = shape.table.cell(0, 0)._tc
    cell_elem.set("gridSpan", "2")

    rule = TableMergedCellsRule()
    ctx = AuditContext(source_name="merged_table.pptx")
    findings = rule.check(prs, ctx)

    assert len(findings) == 1
    f = findings[0]
    assert f.rule_id == "table-merged-cells"
    assert f.sc == "1.3.1"
    assert f.severity == "serious"
    assert f.why_unfixable is not None
    assert len(f.manual_steps) >= 3
