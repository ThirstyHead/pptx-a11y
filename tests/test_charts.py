"""Tests for chart and embedded object alt text rule (WCAG 1.1.1 / 4.1.2)."""
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches
from pptx_a11y.rules import ChartAltTextRule, AuditContext


def test_chart_missing_alt_detected():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Simulate chart graphic frame
    tx = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(3))
    tx.name = "Revenue Chart"
    chart_elem = parse_xml('<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId2"/>')
    tx._element.spPr.append(chart_elem)

    rule = ChartAltTextRule()
    findings = rule.check(prs, AuditContext("charts.pptx"))
    assert len(findings) == 1
    f = findings[0]
    assert f.rule_id == "chart-missing-alt"
    assert f.sc == "1.1.1"
    assert f.severity == "critical"
    assert f.why_unfixable is not None
    assert len(f.manual_steps) >= 3
