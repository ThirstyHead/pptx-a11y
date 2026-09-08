"""Deterministic remediation engine for PowerPoint presentations."""
from pathlib import Path
from typing import Dict
from pptx import Presentation


def remediate_presentation(
    in_path: str | Path,
    out_path: str | Path,
    default_lang: str = "en-US",
) -> Dict[str, int]:
    prs = Presentation(str(in_path))
    fixes = {
        "title_added": 0,
        "table_headers_set": 0,
        "language_tagged": 0,
    }

    # 1. Remediate core title if missing
    if not (prs.core_properties.title or "").strip():
        inferred_title = Path(in_path).stem.replace("-", " ").replace("_", " ").title()
        try:
            if prs.slides and prs.slides[0].shapes.title and prs.slides[0].shapes.title.text.strip():
                inferred_title = prs.slides[0].shapes.title.text.strip()
        except Exception:
            pass
        prs.core_properties.title = inferred_title
        fixes["title_added"] += 1

    # 2. Iterate slides for table headers and language tags
    for slide in prs.slides:
        for shape in slide.shapes:
            # Fix table headers
            if shape.has_table:
                tblPr = shape._element.xpath(".//a:tblPr")
                if tblPr and tblPr[0].get("firstRow") not in ("1", "true"):
                    tblPr[0].set("firstRow", "1")
                    fixes["table_headers_set"] += 1

            # Fix language tags on text runs
            tf = getattr(shape, "text_frame", None)
            if tf:
                for p in tf.paragraphs:
                    for run in p.runs:
                        rPr = run._r.xpath(".//a:rPr")
                        if rPr and not rPr[0].get("lang"):
                            rPr[0].set("lang", default_lang)
                            fixes["language_tagged"] += 1
                        elif not rPr:
                            new_rPr = run._r.get_or_add_rPr()
                            new_rPr.set("lang", default_lang)
                            fixes["language_tagged"] += 1

    prs.save(str(out_path))
    return fixes
